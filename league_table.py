"""The paper's OWN ranking metric (Fig 3d: one-sample t over the 16 clock effect sizes
per intervention; MTC cutoff 0.00625 = Bonferroni at Me=8, their Supplementary Table 20),
applied to all 51 arms. Reproduces their published counts exactly: 19 / 13 / 26 / 3."""
import openpyxl, statistics as st
from math import lgamma, log, exp, sqrt
def _bcf(a,b,x):
    MAXIT,EPS,FPMIN=300,3e-16,1e-300
    qab,qap,qam=a+b,a+1,a-1; c,d=1.0,1-qab*x/qap
    d=1/(d if abs(d)>FPMIN else FPMIN); h=d
    for m in range(1,MAXIT+1):
        m2=2*m
        aa=m*(b-m)*x/((qam+m2)*(a+m2)); d=1+aa*d; d=1/(d if abs(d)>FPMIN else FPMIN)
        c=1+aa/c; c=c if abs(c)>FPMIN else FPMIN; h*=d*c
        aa=-(a+m)*(qab+m)*x/((a+m2)*(qap+m2)); d=1+aa*d; d=1/(d if abs(d)>FPMIN else FPMIN)
        c=1+aa/c; c=c if abs(c)>FPMIN else FPMIN; de=d*c; h*=de
        if abs(de-1)<EPS: break
    return h
def _bi(a,b,x):
    if x<=0: return 0.0
    if x>=1: return 1.0
    bt=exp(lgamma(a+b)-lgamma(a)-lgamma(b)+a*log(x)+b*log(1-x))
    return bt*_bcf(a,b,x)/a if x<(a+1)/(a+b+2) else 1-bt*_bcf(b,a,1-x)/b
def tp(t,df): return _bi(df/2,0.5,df/(df+t*t))
wb=openpyxl.load_workbook('SupplementaryTable1_Sehgal2026.xlsx',data_only=True)
rE=[r for r in wb['Effect Sizes'].iter_rows(values_only=True) if r[0] is not None]
h=list(rE[0]); D=rE[1:]
rP=[r for r in wb['Pvalues'].iter_rows(values_only=True) if any(c is not None for c in r)]
hP=list(rP[0]); P={r[0]:r for r in rP[1:]}
CL16=['Horvath1','Horvath2','Hannum','PCHorvath1','PCHorvath2','PCHannum','DNAmEMRAge','OMICmAge',
      'PhenoAge','GrimAgeV1','GrimAgeV2','PCPhenoAge','PCGrimAge','SystemsAge','DunedinPoAm38','DunedinPACE']
CUT=0.00625
CTRL={'Healthy Guidelines Diet','Low Fat Diet','Low Fat Diet + Exercise','Low Fat Diet + no Exercise',
      'Omnivore Diet','HBOT - Mild Pressure','HBT - Mild Pressure'}
res=[]
for r in D:
    v=[r[h.index(c)] for c in CL16]
    m=st.mean(v); s=st.stdev(v); t=m/(s/sqrt(16)); p=tp(t,15)
    tags=[]
    if r[0] in CTRL: tags.append('COMPARATOR/SHAM ARM')
    if str(r[h.index('Condition')]) not in ('None','NA','nan'): tags.append('disease:'+str(r[h.index('Condition')]))
    if r[h.index('Link to main publication (PMID)')] in (None,'None','NA'): tags.append('UNPUBLISHED')
    if str(r[h.index('Study Design: Allocation')])=='Randomized': tags.append('randomized')
    res.append((m,p,r[0],P[r[0]][17],tags))
res.sort(key=lambda x:x[0])
print("THE PAPER'S OWN LEAGUE TABLE OF LONGEVITY INTERVENTIONS")
print("(Fig 3d metric; verdict column uses the paper's own MTC cutoff p<0.00625)\n")
print(f"{'#':>3s} {'intervention':34s} {'n':>4s} {'mean':>8s} {'p':>9s}  {'verdict':22s} notes")
print("-"*140)
for i,(m,p,a,n,tags) in enumerate(res,1):
    if p<CUT: verdict = 'REJUVENATES **' if m<0 else 'ACCELERATES **'
    elif p<0.05: verdict = 'rejuvenates *' if m<0 else 'accelerates *'
    else: verdict = 'no effect'
    print(f"{i:3d} {a[:34]:34s} {n:4.0f} {m:+8.3f} {p:9.4f}  {verdict:22s} {', '.join(tags)}")
print("-"*140)
print("\nWHERE THE BEST-EVIDENCED INTERVENTIONS LAND:")
for key in ['Exercise','CR / IMF','Diet + Exercise','Vegan Diet','Semaglutide','Gastric Bypass',
            'Rapamycin dose 1','Rapamycin dose 3','Metformin','Green Mediterranean Diet',
            'AC11 Supplement','Lavendar Oil Supplement','Botanical Supplement','Healthy Guidelines Diet',
            'Low Fat Diet','Omnivore Diet','DQ Senolytics 1','Follistatin Gene Therapy']:
    for i,(m,p,a,n,tags) in enumerate(res,1):
        if a==key:
            v = ('REJUVENATES**' if m<0 else 'ACCELERATES**') if p<CUT else (('rejuv*' if m<0 else 'accel*') if p<0.05 else 'no effect')
            print(f"   rank {i:2d}/51  {a[:30]:32s} mean={m:+.3f} p={p:.4f}  -> {v}")
