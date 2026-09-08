"""Randomized between-arm contrasts, NON-NESTED pairs only, with 95% CIs.
Committed so every number in the letter is reproducible."""
import openpyxl, statistics as st
from math import lgamma, log, exp, sqrt, erfc
def _bcf(a,b,x):
    MAXIT,EPS,FPMIN=300,3e-16,1e-300
    qab,qap,qam=a+b,a+1,a-1; c,d=1.0,1-qab*x/qap
    d=1/(d if abs(d)>FPMIN else FPMIN); h=d
    for m in range(1,MAXIT+1):
        m2=2*m
        aa=m*(b-m)*x/((qam+m2)*(a+m2)); d=1+aa*d; d=1/(d if abs(d)>FPMIN else FPMIN)
        c=1+aa/c; c=c if abs(c)>FPMIN else FPMIN; h*=d*c
        aa=-(a+m)*(qab+m)*x/((a+m2)*(qap+m2)); d=1+aa*d; d=1/(d if abs(d)>FPMIN else FPMIN)
        c=1+aa/c; c=c if abs(c)>FPMIN else FPMIN; de=d*c; h*=de
        if abs(de-1)<EPS: break
    return h
def _bi(a,b,x):
    if x<=0: return 0.0
    if x>=1: return 1.0
    bt=exp(lgamma(a+b)-lgamma(a)-lgamma(b)+a*log(x)+b*log(1-x))
    return bt*_bcf(a,b,x)/a if x<(a+1)/(a+b+2) else 1-bt*_bcf(b,a,1-x)/b
def tp(t,df): return _bi(df/2,0.5,df/(df+t*t))
def t_from_p(p,df):
    lo,hi=0.0,400.0
    for _ in range(300):
        mid=(lo+hi)/2; lo,hi=(mid,hi) if tp(mid,df)>p else (lo,mid)
    return (lo+hi)/2
zp=lambda z: erfc(abs(z)/sqrt(2))
wb=openpyxl.load_workbook('SupplementaryTable1_Sehgal2026.xlsx',data_only=True)
rE=[r for r in wb['Effect Sizes'].iter_rows(values_only=True) if r[0] is not None]
h=list(rE[0]); E={r[0]:r for r in rE[1:]}
rP=[r for r in wb['Pvalues'].iter_rows(values_only=True) if any(c is not None for c in r)]
hP=list(rP[0]); P={r[0]:r for r in rP[1:]}
eff=lambda a,c: E[a][h.index(c)]; pv=lambda a,c: P[a][hP.index(c)]; nn=lambda a: P[a][17]
def se(a,c):
    t=t_from_p(pv(a,c),nn(a)-1); return abs(eff(a,c))/t if t>1e-9 else None
# Paper's published pooled single-arm estimates (Supplementary Table 7, n=51)
PUB={'GrimAgeV2':-0.0814,'PCPhenoAge':-0.08876,'PCGrimAge':-0.07843,'SystemsAge':-0.05859,'DunedinPACE':-0.0891}
HEAD=list(PUB)
# NON-NESTED pairs only: one contrast per trial, no shared participants, no shared comparator
PAIRS=[('DIRECT-PLUS','Green Mediterranean Diet','Healthy Guidelines Diet'),
       ('CENTRAL','Low Carb Diet','Low Fat Diet'),
       ('Twin vegan','Vegan Diet','Omnivore Diet'),
       ('HBOT vs sham','HBOT - High Pressure','HBOT - Mild Pressure'),
       ('HBT vs sham','HBT - High Pressure','HBT - Mild Pressure')]
for label,SET in [('5 non-nested contrasts',PAIRS),('4 contrasts, HBT (n=5 vs 2) dropped',PAIRS[:4])]:
    print('='*94); print(label); print('='*94)
    print(f"{'clock':13s} {'randomized contrast':>21s} {'95% CI':>22s} {'P':>7s} | {'paper pooled':>13s} {'excluded?':>10s}")
    for c in HEAD:
        arr=[]
        for _,t_,ct in SET:
            s1,s2=se(t_,c),se(ct,c)
            if s1 and s2: arr.append((eff(t_,c)-eff(ct,c), sqrt(s1**2+s2**2)))
        w=[1/s**2 for _,s in arr]
        est=sum(wi*d for wi,(d,_) in zip(w,arr))/sum(w); sep=sqrt(1/sum(w))
        lo,hi=est-1.96*sep, est+1.96*sep
        excl = "YES" if not (lo <= PUB[c] <= hi) else "no"
        print(f"{c:13s} {est:+21.4f} [{lo:+.3f}, {hi:+.3f}]{'':>4s} {zp(est/sep):7.3f} | {PUB[c]:+13.4f} {excl:>10s}")
    print()
